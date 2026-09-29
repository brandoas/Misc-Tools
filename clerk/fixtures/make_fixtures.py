"""Generate the ITS 3100 Deliverable 01 fixture reports and their answer key.

There is no archive of graded student reports to calibrate against, so ground
truth is built instead of found. Each fixture is written to satisfy or to fail
specific Evidence bullets on cards/3100-d01.md, and expected.json records the
rating that follows.

That gives a regression suite. Change the card, change the model, change a
prompt, rerun, and see exactly what moved.

The limit of this method, stated plainly: these are written by a machine and are
cleaner, better organised, and more consistently formatted than real student
work. They test the card, the plumbing, and the rating rule. They do not tell
you how the judge behaves on a real submission that wanders. The first real
deliverable is still an iteration.

    python fixtures/make_fixtures.py
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).parent
RUN = HERE.parent / "runs" / "3100-d01-fixtures"

# One student's coordinated allocation, used consistently across the fixtures so
# a report reads like one network rather than a pile of examples.
LINK_LOCAL = "fe80::e6d:3cff:fe26:0"
WAN_SLAAC = "2610:a8:483f:c00:e6d:3cff:fe26:0/64"
WAN_DHCPV6 = "2610:a8:4831:38e::9d/128"
LAN1_ADDR = "2610:a8:483f:c1c::1/64"
LAN2_ADDR = "2610:a8:483f:c1d::1/64"
CLIENT_ADDR = "2610:a8:483f:c1c:ad1f:1db:8b7:837d"
GATEWAY_V4 = "172.22.16.254"
UBUNTU_V4 = "172.22.16.10"

PING_GW = f"""\
brandon@ubuntu-gui-1:~$ ping -c 3 {GATEWAY_V4}
PING {GATEWAY_V4} ({GATEWAY_V4}) 56(84) bytes of data.
64 bytes from {GATEWAY_V4}: icmp_seq=1 ttl=64 time=0.612 ms
64 bytes from {GATEWAY_V4}: icmp_seq=2 ttl=64 time=0.548 ms
64 bytes from {GATEWAY_V4}: icmp_seq=3 ttl=64 time=0.531 ms

--- {GATEWAY_V4} ping statistics ---
3 packets transmitted, 3 received, 0% packet loss, time 2040ms"""

PING_8888 = """\
brandon@ubuntu-gui-1:~$ ping -c 3 8.8.8.8
PING 8.8.8.8 (8.8.8.8) 56(84) bytes of data.
64 bytes from 8.8.8.8: icmp_seq=1 ttl=113 time=14.2 ms
64 bytes from 8.8.8.8: icmp_seq=2 ttl=113 time=13.9 ms
64 bytes from 8.8.8.8: icmp_seq=3 ttl=113 time=14.1 ms

--- 8.8.8.8 ping statistics ---
3 packets transmitted, 3 received, 0% packet loss, time 2003ms"""

SHOW_INT_WAN = f"""\
vyos@VyOS-1:~$ show interfaces
Codes: S - State, L - Link, u - Up, D - Down, A - Admin Down
Interface    IP Address                              S/L  Description
---------    ----------                              ---  -----------
eth0         {WAN_SLAAC}   u/u  WAN
             {WAN_DHCPV6}
             {LINK_LOCAL}/64"""

PING6 = """\
vyos@VyOS-1:~$ ping6 2606:4700:4700::1111 count 3
PING 2606:4700:4700::1111 (2606:4700:4700::1111) 56 data bytes
64 bytes from 2606:4700:4700::1111: icmp_seq=1 ttl=57 time=11.8 ms
64 bytes from 2606:4700:4700::1111: icmp_seq=2 ttl=57 time=11.4 ms
64 bytes from 2606:4700:4700::1111: icmp_seq=3 ttl=57 time=11.6 ms

--- 2606:4700:4700::1111 ping statistics ---
3 packets transmitted, 3 received, 0% packet loss, time 2003ms"""

SHOW_ROUTE = """\
vyos@VyOS-1:~$ show ipv6 route
Codes: K - kernel route, C - connected, S - static, R - RIPng,
       O - OSPFv3, I - IS-IS, B - BGP, * - FIB route
S>* ::/0 [1/0] via fe80::200:5eff:fe00:101, eth0, weight 1, 00:04:12
C>* 2610:a8:483f:c00::/64 is directly connected, eth0, 00:04:12
C>* fe80::/64 is directly connected, eth0, 00:06:30"""

SHOW_INT_PD = f"""\
vyos@VyOS-1:~$ show interfaces
Interface    IP Address                              S/L  Description
---------    ----------                              ---  -----------
eth0         {WAN_SLAAC}   u/u  WAN
             {WAN_DHCPV6}
eth1         172.22.16.254/24                        u/u  LAN1
             {LAN1_ADDR}
eth2         172.22.32.254/24                        u/u  LAN2
             {LAN2_ADDR}"""

PD_CONFIG = """\
vyos@VyOS-1:~$ show configuration commands | match dhcpv6-options
set interfaces ethernet eth0 dhcpv6-options pd 0 length '62'
set interfaces ethernet eth0 dhcpv6-options pd 0 interface eth1 address '1'
set interfaces ethernet eth0 dhcpv6-options pd 0 interface eth1 sla-id '0'
set interfaces ethernet eth0 dhcpv6-options pd 0 interface eth2 address '1'
set interfaces ethernet eth0 dhcpv6-options pd 0 interface eth2 sla-id '1'"""

NAME_SERVER_BEFORE = """\
vyos@VyOS-1:~$ show configuration commands | match name-server
set system name-server '132.235.6.1'
set system name-server '132.235.6.2'"""

NAME_SERVER_AFTER = """\
vyos@VyOS-1:~$ show configuration commands | match name-server
set system name-server '2610:a8:4831:300::10'
set system name-server '2610:a8:4831:300::11'"""

PING_GOOGLE = """\
vyos@VyOS-1:~$ ping google.com count 2
PING google.com(iad23s60-in-x0e.1e100.net (2607:f8b0:4004:c07::8e)) 56 data bytes
64 bytes from iad23s60-in-x0e.1e100.net (2607:f8b0:4004:c07::8e): icmp_seq=1 ttl=115 time=13.1 ms
64 bytes from iad23s60-in-x0e.1e100.net (2607:f8b0:4004:c07::8e): icmp_seq=2 ttl=115 time=12.8 ms"""

RA_CONFIG = """\
vyos@VyOS-1:~$ show configuration commands | match router-advert
set service router-advert interface eth1 prefix '::/64'
set service router-advert interface eth1 name-server '2610:a8:4831:300::10'
set service router-advert interface eth1 name-server '2610:a8:4831:300::11'"""

RA_CAPTURE = """\
No.  Time        Source                    Destination  Protocol  Info
  4  2.104882    fe80::e6d:3cff:fe26:0     ff02::1      ICMPv6    Router Advertisement from 0c:6d:3c:26:00:00
 12  8.106114    fe80::e6d:3cff:fe26:0     ff02::1      ICMPv6    Router Advertisement from 0c:6d:3c:26:00:00

Frame 4: Router Advertisement
    Prefix Information: 2610:a8:483f:c1c::/64
        Flag: Autonomous Address Configuration: Set
        Valid Lifetime: 2592000
    Recursive DNS Server: 2610:a8:4831:300::10
    Recursive DNS Server: 2610:a8:4831:300::11"""

NMCLI = f"""\
brandon@ubuntu-gui-1:~$ nmcli -g IP6.ADDRESS device show ens3
{CLIENT_ADDR}/64 | 2610:a8:483f:c1c:5f9a:2e11:ba37:1c02/64 | {LINK_LOCAL.replace('e6d', 'a1f')}/64"""

DIG_AT = """\
brandon@ubuntu-gui-1:~$ dig www.ohio.edu @2610:a8:4831:300::10

;; ANSWER SECTION:
www.ohio.edu.        300    IN    A    132.235.44.62

;; Query time: 4 msec
;; SERVER: 2610:a8:4831:300::10#53(2610:a8:4831:300::10) (UDP)
;; WHEN: Fri Sep 26 09:12:44 EDT 2026"""

DIG_NO_AT = """\
brandon@ubuntu-gui-1:~$ dig www.ohio.edu

;; ANSWER SECTION:
www.ohio.edu.        291    IN    A    132.235.44.62

;; Query time: 0 msec
;; SERVER: 127.0.0.53#53(127.0.0.53) (UDP)"""

NAT64_CONFIG = """\
vyos@VyOS-1:~$ show configuration commands | match nat64
set nat64 source rule 1 source prefix '64:ff9b::/96'

vyos@VyOS-1:~$ show configuration commands | match 'router-advert interface eth2'
set service router-advert interface eth2 prefix '::/64'
set service router-advert interface eth2 name-server '2001:4860:4860::6464'"""

WIN_PING_OHIO = """\
C:\\Users\\Student>ping www.ohio.edu

Pinging www.ohio.edu [64:ff9b::84eb:2c3e] with 32 bytes of data:
Reply from 64:ff9b::84eb:2c3e: time=15ms
Reply from 64:ff9b::84eb:2c3e: time=14ms
Reply from 64:ff9b::84eb:2c3e: time=14ms"""

WIN_PING_8888 = """\
C:\\Users\\Student>ping 8.8.8.8

Pinging 8.8.8.8 with 32 bytes of data:
PING: transmit failed. General failure.
PING: transmit failed. General failure.

Ping statistics for 8.8.8.8:
    Packets: Sent = 2, Received = 0, Lost = 2 (100% loss)"""

CAP_LAN2 = """\
Capture on VyOS-1 eth2 (LAN2, IPv6 only)
No.  Time      Source                        Destination            Protocol  Info
  7  1.042113  2610:a8:483f:c1d:1c4f:9e2:31  64:ff9b::84eb:2c3e     ICMPv6    Echo (ping) request id=0x0001 seq=1
  9  1.057440  64:ff9b::84eb:2c3e            2610:a8:483f:c1d:1c4f:9e2:31  ICMPv6  Echo (ping) reply id=0x0001 seq=1"""

CAP_WAN = """\
Capture on VyOS-1 eth0 (WAN, after translation)
No.  Time      Source            Destination      Protocol  Info
  8  1.048902  132.235.207.157   132.235.44.62    ICMP      Echo (ping) request id=0x0001 seq=1
 10  1.055118  132.235.44.62     132.235.207.157  ICMP      Echo (ping) reply id=0x0001 seq=1"""


def para(doc, text, *, mono=False, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if mono:
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    return p


def block(doc, text):
    """A fixed width block, as the report guidelines require for CLI output."""
    for line in text.split("\n"):
        para(doc, line, mono=True)


def screenshot(doc, text, path: Path):
    """Render the same text as an image, which is a guidelines violation.

    Used by the screenshot-heavy fixture. The content is correct, so a judge
    that only reads text will score it Missing where it should be at most a
    formatting deduction. That gap is exactly what the vision stage is for, and
    this fixture is how the gap gets measured.
    """
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype("consola.ttf", 15)
    except OSError:
        font = ImageFont.load_default()
    lines = text.split("\n")
    width = max(760, max(len(l) for l in lines) * 9 + 40)
    height = len(lines) * 20 + 30
    img = Image.new("RGB", (width, height), (12, 12, 12))
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((14, 12 + i * 20), line, fill=(220, 220, 220), font=font)
    img.save(path)
    doc.add_picture(str(path))


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


# Analysis prose at three depths. The difference between these is the whole
# point of the exercise: "thin" is fluent, correct, and satisfies almost
# nothing, because it restates the configuration instead of analysing the data.

ANALYSIS_T3_FULL = f"""\
The eth0 interface carries three addresses, each acquired a different way.

{LINK_LOCAL}/64 is the link local address. It sits in fe80::/10 and was formed \
automatically when the interface came up. The interface ID e6d:3cff:fe26:0 is \
EUI-64 derived, built from the MAC 0c:6d:3c:26:00:00 by inserting fffe in the \
middle and flipping the universal/local bit.

{WAN_SLAAC} is the SLAAC address, produced by the autoconf command in step 5. \
The prefix 2610:a8:483f:c00::/64 came from the router advertisement sent by the \
upstream ISP router, and VyOS supplied the same EUI-64 interface ID. No server \
was involved: the router advertised a prefix and the host built its own address.

{WAN_DHCPV6} is the DHCPv6 address from step 8. It was assigned by a DHCPv6 \
server, not derived, which is why it carries a /128 rather than a /64: the \
server handed out exactly one address and no on-link prefix.

The distinction matters. SLAAC is stateless and the host chooses its own \
interface ID, while DHCPv6 is stateful and the server records what it gave out."""

ANALYSIS_T3_THIN = """\
After running the autoconf and dhcpv6 commands the interface has several \
addresses. The show interfaces command displays them all. IPv6 allows multiple \
addresses on an interface, which is one of its advantages over IPv4. The \
configuration worked as expected and connectivity was confirmed."""

ANALYSIS_T4_FULL = f"""\
The delegated prefix is 2610:a8:483f:c1c::/62, handed down by the ISP in \
response to the prefix delegation request configured in step 11. A /62 provides \
four /64 subnets, and the sla-id selects which one each interface takes.

eth1 has sla-id 0, so it takes the first /64 of the delegation and its address \
is {LAN1_ADDR}. eth2 has sla-id 1, so it takes the second and its address is \
{LAN2_ADDR}. The difference between c1c and c1d in those two addresses is the \
sla-id being added to the base of the delegated prefix.

The address suffix ::1 in both comes from the "address 1" line, which tells \
VyOS to take host 1 of whichever subnet the sla-id selected.

This is the link between the WAN and the LANs: nothing on the LAN interfaces \
was configured by hand. The WAN interface requested a block, the ISP delegated \
it, and VyOS carved it into per interface /64s automatically."""

ANALYSIS_T4_THIN = """\
The DHCP-PD configuration was applied to eth0 and the LAN interfaces received \
addresses. The show interfaces command shows the new addresses on eth1 and \
eth2. Prefix delegation is a useful feature that makes IPv6 deployment easier \
on internal networks."""

ANALYSIS_T5_FULL = """\
Before the change the router used 132.235.6.1 and 132.235.6.2, which are IPv4 \
addresses, yet the ping to google.com returned an IPv6 address \
(2607:f8b0:4004:c07::8e). That is the distinction between DNS transport and DNS \
record type. The query travelled to the name server over IPv4, but the record \
it returned was a AAAA record carrying an IPv6 address. The two are independent: \
an IPv4 resolver can serve AAAA records, and an IPv6 resolver can serve A \
records.

After deleting the old name servers and setting the ITS IPv6 resolvers, both \
the transport and the records are IPv6. The delete was necessary because set \
adds to the list rather than replacing it, so without it the router would have \
kept four name servers and preferred the original two."""

ANALYSIS_T5_THIN = """\
DNS was reconfigured to use the IPv6 name servers as instructed. The name \
server setting is shown with show configuration commands. Ping to google.com \
worked both before and after the change."""

ANALYSIS_T6_FULL = f"""\
The capture shows router advertisements leaving the VyOS eth1 interface every \
six seconds, sent to the all nodes multicast address ff02::1. Frame 4 carries \
the prefix information option with 2610:a8:483f:c1c::/64 and the Autonomous \
Address Configuration flag set, which is what tells a listening host it may \
build its own address from that prefix. The same frame carries both recursive \
DNS server options.

The Ubuntu client address from nmcli is {CLIENT_ADDR}. Its first 64 bits are \
exactly the prefix advertised in frame 4, so the network portion came from the \
router advertisement. The interface ID ad1f:1db:8b7:837d is not EUI-64: the \
MAC is 0c:6d:3c:26:00:00 and an EUI-64 interface ID would read \
e6d:3cff:fe26:0. It is an RFC 7217 stable privacy address, which \
NetworkManager generates by hashing the prefix with a host secret. The second \
address in the nmcli output is a rotating temporary address.

The dig with the explicit @ server returns an A record, 132.235.44.62, and the \
SERVER line reads 2610:a8:4831:300::10, naming the IPv6 address that actually \
answered. This is an IPv4 record delivered over IPv6 transport. Running dig \
without the @ gives the same answer but the SERVER line reads 127.0.0.53, the \
local systemd stub resolver, which hides which upstream server did the work."""

ANALYSIS_T6_THIN = """\
Router advertisements were configured on eth1 and the capture shows them being \
sent. The Ubuntu client picked up an address automatically. The dig command \
returned the correct answer for www.ohio.edu."""

ANALYSIS_T7_FULL = """\
The capture on LAN2 shows the Windows host sending an ICMPv6 echo request from \
its own address to 64:ff9b::84eb:2c3e. That destination is the NAT64 well known \
prefix 64:ff9b::/96 with the IPv4 address of www.ohio.edu embedded in the last \
32 bits: 84eb:2c3e in hexadecimal is 132.235.44.62 in dotted decimal \
(0x84=132, 0xeb=235, 0x2c=44, 0x3e=62). DNS64 synthesised that AAAA record \
because www.ohio.edu has no real AAAA.

The capture on the WAN shows the same ping after translation. It is now ICMP \
over IPv4, sourced from the router's own WAN address 132.235.207.157 and \
destined for 132.235.44.62, with the same ICMP identifier 0x0001 and sequence 1 \
so the two captures can be matched to each other. The reply takes the reverse \
path: it arrives as IPv4 on the WAN and leaves LAN2 as ICMPv6 with the \
64:ff9b:: source restored.

The ping to 8.8.8.8 fails, and the reason is structural rather than a \
misconfiguration. A literal IPv4 address typed into ping never goes through \
DNS at all, so DNS64 never sees it and never synthesises a 64:ff9b:: address \
for it. With no IPv6 destination, an IPv6 only host has nothing to send to.

464XLAT closes that gap. The client runs a CLAT, a stateless translator that \
presents a fake IPv4 interface to applications and translates their IPv4 \
packets into IPv6 destined for the NAT64 prefix, which the PLAT at the network \
edge then translates back to IPv4. The CLAT has to know which prefix to use, \
and the PREF64 option in the router advertisement (RFC 8781) is how the network \
tells it, rather than making the client probe for it. So a literal IPv4 \
address works again, because the translation happens on the host before DNS \
would ever have been consulted.

Reference: RFC 8781, Discovering PREF64 in Router Advertisements, \
https://www.rfc-editor.org/rfc/rfc8781.html"""

ANALYSIS_T7_THIN = """\
NAT64 was enabled and the Windows machine was able to ping www.ohio.edu. The \
captures were taken on both interfaces. Pinging 8.8.8.8 did not work. 464XLAT \
is a technology that allows IPv4 traffic over IPv6 only networks and is widely \
used by mobile carriers."""

AQ1_FULL = f"""\
My eth0 interface carried three addresses at once: the link local \
{LINK_LOCAL}, the SLAAC address {WAN_SLAAC}, and the DHCPv6 address \
{WAN_DHCPV6}. The Ubuntu client carried three as well: a stable privacy \
address, a rotating temporary address, and its own link local.

Having several is normal in IPv6, so the stack needs a rule for picking one. \
RFC 6724 defines it as an ordered set of rules applied to candidate source \
addresses for a given destination. Scope is matched first, so a link local \
destination gets a link local source and my {LINK_LOCAL} would never be chosen \
to reach the internet. Then temporary addresses are preferred over stable ones \
for outbound connections, which is why my Ubuntu client would source a new \
connection from its rotating address rather than from {CLIENT_ADDR}. Where \
several candidates remain, the one sharing the longest prefix with the \
destination wins, so reaching another host on my LAN1 would select the \
{LAN1_ADDR.split('/')[0]} side rather than the WAN address."""

AQ1_THIN = """\
IPv6 interfaces often have more than one address. The operating system picks \
one of them when it sends a packet. There are rules that govern this choice and \
they are defined in an RFC."""

AQ2_FULL = """\
The workflow starts at the customer router, which sends a DHCPv6 Solicit \
carrying an IA_PD option. That option is the request for a prefix rather than \
for a single address. The ISP's DHCPv6 server replies with an Advertise, then \
the router sends a Request and the server confirms with a Reply that carries \
the delegated prefix, in my case a /62. The router then subdivides that block \
and assigns a /64 out of it to each internal interface, using the sla-id to \
choose which one, and begins advertising those /64s on its LANs so that clients \
can build their own addresses by SLAAC. The lease is renewed like any other \
DHCPv6 binding.

This matters for residential and small business adoption because it makes the \
internal network self configuring. Without it, every site would need a static \
allocation arranged with the ISP and configured by hand on the router, which is \
not something a home user or a small office without an administrator will ever \
do. Prefix delegation means the router asks and the ISP answers, and the \
customer never sees an address.

It also removes the reason NAT existed. In IPv4 a site received one public \
address and NAT was the only way to put more than one device behind it, which \
cost end to end reachability and made inbound connections a configuration \
problem. With DHCP-PD every device at the site gets a genuinely globally \
routable address from a block the ISP routes to the site, so there is nothing \
for NAT to solve. NAT in IPv6 is not a security feature being given up, since \
that role belongs to the firewall either way."""

AQ2_THIN = """\
The router asks the ISP for a prefix and the ISP gives it one. The router then \
uses that prefix on its internal networks. This is important because it makes \
IPv6 easier to deploy at home and in small businesses, and it means NAT is not \
needed."""

AQ3_FULL = """\
SLAAC appears twice in this lab on opposite sides of the protocol.

In Task 3 VyOS-1 is the SLAAC client. The autoconf command on eth0 made the \
router listen for router advertisements from the upstream ISP router and build \
its own WAN address from the prefix it heard. It was receiving.

In Task 6 VyOS-1 is the SLAAC server, or more precisely the advertising router. \
The router-advert configuration on eth1 made it originate router \
advertisements onto LAN1, carrying the prefix and the DNS servers, so that the \
Ubuntu client could do on LAN1 what the router itself had done on the WAN. It \
was sending.

The difference is direction and role. The same protocol, the same message type, \
but in Task 3 the router consumes an advertisement and in Task 6 it produces \
one."""

AQ3_THIN = """\
SLAAC was used in both Task 3 and Task 6. It allows hosts to configure their \
own addresses automatically without a DHCP server."""


def build_strong(doc, imgdir):
    heading(doc, "ITS 3100 Deliverable 01 - IPv6")
    para(doc, "Brandon Saunders, ITS 3100, Fall 2026")

    heading(doc, "Task 2 - Deploy", 2)
    para(doc, "The network was built to match the physical diagram in the "
              "handout: VyOS-1 eth0 is the WAN, eth1 is LAN1 carrying "
              "Ubuntu-GUI, and eth2 is LAN2 carrying the Windows client. The "
              "diagram below was drawn in Draw.io and carries object names, "
              "network numbers and interface addresses.")
    doc.add_picture(str(imgdir / "diagram.png"))
    para(doc, f"Ubuntu-GUI was configured with the static address {UBUNTU_V4}/24 "
              f"and gateway {GATEWAY_V4}. No DNS server was configured at this "
              f"stage, as the handout requires.")
    block(doc, PING_GW)
    block(doc, PING_8888)
    para(doc, "Both pings succeed, so the IPv4 network is functional out to the "
              "internet before any IPv6 configuration is added.")

    heading(doc, "Task 3 - IPv6 WAN configuration with SLAAC", 2)
    block(doc, SHOW_INT_WAN)
    block(doc, PING6)
    block(doc, SHOW_ROUTE)
    para(doc, ANALYSIS_T3_FULL)

    heading(doc, "Task 4 - DHCP-PD on the LAN interfaces", 2)
    block(doc, PD_CONFIG)
    block(doc, SHOW_INT_PD)
    para(doc, ANALYSIS_T4_FULL)

    heading(doc, "Task 5 - IPv6 and DNS", 2)
    block(doc, NAME_SERVER_BEFORE)
    block(doc, PING_GOOGLE)
    block(doc, NAME_SERVER_AFTER)
    para(doc, ANALYSIS_T5_FULL)

    heading(doc, "Task 6 - Router advertisements on LAN1", 2)
    block(doc, RA_CONFIG)
    block(doc, RA_CAPTURE)
    block(doc, NMCLI)
    block(doc, DIG_AT)
    block(doc, DIG_NO_AT)
    para(doc, ANALYSIS_T6_FULL)

    heading(doc, "Task 7 - IPv6 only network and NAT64", 2)
    block(doc, NAT64_CONFIG)
    block(doc, WIN_PING_OHIO)
    block(doc, CAP_LAN2)
    block(doc, CAP_WAN)
    block(doc, WIN_PING_8888)
    para(doc, ANALYSIS_T7_FULL)

    heading(doc, "Additional Questions", 2)
    heading(doc, "Question 1 - Source address selection", 3)
    para(doc, AQ1_FULL)
    heading(doc, "Question 2 - DHCP-PD workflow", 3)
    para(doc, AQ2_FULL)
    heading(doc, "Question 3 - SLAAC in Tasks 3 and 6", 3)
    para(doc, AQ3_FULL)


def build_thin(doc, imgdir):
    """Everything present, nothing analysed. The failure mode that matters."""
    heading(doc, "ITS 3100 Deliverable 01 - IPv6")
    para(doc, "Jordan Ellis, ITS 3100, Fall 2026")

    heading(doc, "Task 2 - Deploy", 2)
    doc.add_picture(str(imgdir / "diagram.png"))
    para(doc, f"Ubuntu-GUI was given the address {UBUNTU_V4}/24 with gateway "
              f"{GATEWAY_V4} and no DNS.")
    block(doc, PING_GW)
    block(doc, PING_8888)

    heading(doc, "Task 3 - IPv6 WAN configuration with SLAAC", 2)
    block(doc, SHOW_INT_WAN)
    block(doc, PING6)
    block(doc, SHOW_ROUTE)
    para(doc, ANALYSIS_T3_THIN)

    heading(doc, "Task 4 - DHCP-PD on the LAN interfaces", 2)
    block(doc, PD_CONFIG)
    block(doc, SHOW_INT_PD)
    para(doc, ANALYSIS_T4_THIN)

    heading(doc, "Task 5 - IPv6 and DNS", 2)
    block(doc, NAME_SERVER_AFTER)
    block(doc, PING_GOOGLE)
    para(doc, ANALYSIS_T5_THIN)

    heading(doc, "Task 6 - Router advertisements on LAN1", 2)
    block(doc, RA_CONFIG)
    block(doc, RA_CAPTURE)
    block(doc, NMCLI)
    block(doc, DIG_AT)
    para(doc, ANALYSIS_T6_THIN)

    heading(doc, "Task 7 - IPv6 only network and NAT64", 2)
    block(doc, NAT64_CONFIG)
    block(doc, WIN_PING_OHIO)
    block(doc, CAP_LAN2)
    block(doc, WIN_PING_8888)
    para(doc, ANALYSIS_T7_THIN)

    heading(doc, "Additional Questions", 2)
    para(doc, AQ1_THIN)
    para(doc, AQ2_THIN)
    para(doc, AQ3_THIN)


def build_partial(doc, imgdir):
    """Specific planted gaps, each one a Critique bullet from the card."""
    heading(doc, "ITS 3100 Deliverable 01 - IPv6")
    para(doc, "Sam Whitaker, ITS 3100, Fall 2026")

    heading(doc, "Task 2 - Deploy", 2)
    doc.add_picture(str(imgdir / "diagram.png"))
    para(doc, f"Ubuntu-GUI was configured with {UBUNTU_V4}/24, gateway "
              f"{GATEWAY_V4}, and no DNS server.")
    block(doc, PING_GW)
    # Planted: only one of the two pings shown.

    heading(doc, "Task 3 - IPv6 WAN configuration with SLAAC", 2)
    # Planted: only the SLAAC address, no DHCPv6 address.
    block(doc, f"""vyos@VyOS-1:~$ show interfaces
Interface    IP Address                              S/L  Description
---------    ----------                              ---  -----------
eth0         {WAN_SLAAC}   u/u  WAN
             {LINK_LOCAL}/64""")
    block(doc, PING6)
    block(doc, SHOW_ROUTE)
    para(doc, "The address on eth0 was acquired using SLAAC after the autoconf "
              "command was applied. The prefix 2610:a8:483f:c00::/64 came from "
              "the upstream router advertisement, and the interface ID "
              "e6d:3cff:fe26:0 is EUI-64 derived from the MAC address.")

    heading(doc, "Task 4 - DHCP-PD on the LAN interfaces", 2)
    block(doc, PD_CONFIG)
    block(doc, SHOW_INT_PD)
    para(doc, ANALYSIS_T4_FULL)

    heading(doc, "Task 5 - IPv6 and DNS", 2)
    block(doc, NAME_SERVER_BEFORE)
    block(doc, PING_GOOGLE)
    block(doc, NAME_SERVER_AFTER)
    para(doc, ANALYSIS_T5_FULL)

    heading(doc, "Task 6 - Router advertisements on LAN1", 2)
    block(doc, RA_CONFIG)
    block(doc, RA_CAPTURE)
    block(doc, NMCLI)
    block(doc, DIG_AT)
    block(doc, DIG_NO_AT)
    # Planted: capture present but never tied to the client address, and the
    # EUI-64 versus privacy distinction is not made.
    para(doc, "The router advertisements are visible in the capture and the "
              "Ubuntu client received an address. The dig command returned an A "
              "record for www.ohio.edu over the IPv6 transport.")

    heading(doc, "Task 7 - IPv6 only network and NAT64", 2)
    block(doc, NAT64_CONFIG)
    block(doc, WIN_PING_OHIO)
    # Planted: only one side of the translation captured.
    block(doc, CAP_LAN2)
    block(doc, WIN_PING_8888)
    para(doc, "The ping to www.ohio.edu is translated and gets a reply, while "
              "the ping to 8.8.8.8 fails because there is no route for IPv4 on "
              "an IPv6 only network. 464XLAT is a combination of a CLAT on the "
              "client and a PLAT in the network that together allow IPv4 "
              "applications to work over an IPv6 only path.")
    # Planted: no reference cited, DNS64 reason not given.

    heading(doc, "Additional Questions", 2)
    # Planted: generic answer with no lab data.
    para(doc, AQ1_THIN)
    para(doc, AQ2_FULL)
    para(doc, AQ3_FULL)


def build_screenshots(doc, imgdir):
    """Correct content, but the CLI output is pasted as images.

    Text only judging should score this badly. The vision stage is what should
    recover it, so this fixture is the measurement of what the vision stage is
    worth.
    """
    heading(doc, "ITS 3100 Deliverable 01 - IPv6")
    para(doc, "Riley Chen, ITS 3100, Fall 2026")

    heading(doc, "Task 2 - Deploy", 2)
    doc.add_picture(str(imgdir / "diagram.png"))
    screenshot(doc, PING_GW, imgdir / "sc-ping-gw.png")
    screenshot(doc, PING_8888, imgdir / "sc-ping-8888.png")
    para(doc, f"Ubuntu-GUI uses {UBUNTU_V4}/24 with gateway {GATEWAY_V4} and no "
              f"DNS configured.")

    heading(doc, "Task 3 - IPv6 WAN configuration with SLAAC", 2)
    screenshot(doc, SHOW_INT_WAN, imgdir / "sc-show-int.png")
    screenshot(doc, PING6, imgdir / "sc-ping6.png")
    screenshot(doc, SHOW_ROUTE, imgdir / "sc-route.png")
    para(doc, ANALYSIS_T3_FULL)

    heading(doc, "Task 4 - DHCP-PD on the LAN interfaces", 2)
    screenshot(doc, PD_CONFIG, imgdir / "sc-pd.png")
    screenshot(doc, SHOW_INT_PD, imgdir / "sc-show-pd.png")
    para(doc, ANALYSIS_T4_FULL)

    heading(doc, "Task 5 - IPv6 and DNS", 2)
    screenshot(doc, NAME_SERVER_BEFORE, imgdir / "sc-ns-before.png")
    screenshot(doc, PING_GOOGLE, imgdir / "sc-ping-google.png")
    screenshot(doc, NAME_SERVER_AFTER, imgdir / "sc-ns-after.png")
    para(doc, ANALYSIS_T5_FULL)

    heading(doc, "Task 6 - Router advertisements on LAN1", 2)
    screenshot(doc, RA_CONFIG, imgdir / "sc-ra.png")
    screenshot(doc, RA_CAPTURE, imgdir / "sc-ra-cap.png")
    screenshot(doc, NMCLI, imgdir / "sc-nmcli.png")
    screenshot(doc, DIG_AT, imgdir / "sc-dig-at.png")
    screenshot(doc, DIG_NO_AT, imgdir / "sc-dig.png")
    para(doc, ANALYSIS_T6_FULL)

    heading(doc, "Task 7 - IPv6 only network and NAT64", 2)
    screenshot(doc, NAT64_CONFIG, imgdir / "sc-nat64.png")
    screenshot(doc, WIN_PING_OHIO, imgdir / "sc-win-ohio.png")
    screenshot(doc, CAP_LAN2, imgdir / "sc-cap-lan2.png")
    screenshot(doc, CAP_WAN, imgdir / "sc-cap-wan.png")
    screenshot(doc, WIN_PING_8888, imgdir / "sc-win-8888.png")
    para(doc, ANALYSIS_T7_FULL)

    heading(doc, "Additional Questions", 2)
    para(doc, AQ1_FULL)
    para(doc, AQ2_FULL)
    para(doc, AQ3_FULL)


def build_missing(doc, imgdir):
    """Tasks 5 and 6 absent entirely, the rest thin."""
    heading(doc, "ITS 3100 Deliverable 01 - IPv6")
    para(doc, "Casey Moreno, ITS 3100, Fall 2026")

    heading(doc, "Task 2 - Deploy", 2)
    doc.add_picture(str(imgdir / "diagram.png"))
    para(doc, f"Ubuntu-GUI configured with {UBUNTU_V4}/24 and gateway "
              f"{GATEWAY_V4}, no DNS.")
    block(doc, PING_GW)
    block(doc, PING_8888)

    heading(doc, "Task 3 - IPv6 WAN configuration with SLAAC", 2)
    block(doc, SHOW_INT_WAN)
    block(doc, PING6)
    block(doc, SHOW_ROUTE)
    para(doc, ANALYSIS_T3_FULL)

    heading(doc, "Task 4 - DHCP-PD on the LAN interfaces", 2)
    block(doc, PD_CONFIG)
    block(doc, SHOW_INT_PD)
    para(doc, ANALYSIS_T4_THIN)

    # Tasks 5 and 6 are simply not here.

    heading(doc, "Task 7 - IPv6 only network and NAT64", 2)
    block(doc, NAT64_CONFIG)
    block(doc, WIN_PING_OHIO)
    block(doc, WIN_PING_8888)
    para(doc, ANALYSIS_T7_THIN)

    heading(doc, "Additional Questions", 2)
    para(doc, AQ1_THIN)
    para(doc, AQ2_THIN)


def make_diagram(path: Path) -> None:
    """A drawn network diagram, light background and few colours.

    Deliberately flat and labelled rather than pretty: what the classifier and
    the vision model should see is a diagram with object names, network numbers
    and interface addresses on it, which is what the report guidelines require.
    """
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype("arial.ttf", 13)
        small = ImageFont.truetype("arial.ttf", 11)
    except OSError:
        font = small = ImageFont.load_default()
    img = Image.new("RGB", (880, 420), "white")
    d = ImageDraw.Draw(img)

    def box(x, y, w, h, label, sub=""):
        d.rectangle([x, y, x + w, y + h], outline=(40, 40, 40), width=2)
        d.text((x + 10, y + 8), label, fill=(0, 0, 0), font=font)
        if sub:
            d.text((x + 10, y + 28), sub, fill=(60, 60, 60), font=small)

    box(360, 170, 160, 80, "VyOS-1", "router")
    box(40, 50, 170, 70, "Internet", "ISP / WAN")
    box(40, 290, 190, 80, "Ubuntu-GUI-1", f"{UBUNTU_V4}/24")
    box(650, 290, 190, 80, "Windows-1", "IPv6 only")

    d.line([210, 85, 440, 85, 440, 170], fill=(40, 40, 40), width=2)
    d.text((450, 92), "eth0  WAN  dhcp", fill=(0, 0, 0), font=small)
    d.text((450, 108), WAN_DHCPV6, fill=(0, 0, 0), font=small)

    d.line([360, 230, 230, 230, 230, 300], fill=(40, 40, 40), width=2)
    d.text((240, 200), "eth1  LAN1", fill=(0, 0, 0), font=small)
    d.text((240, 216), f"{GATEWAY_V4}/24", fill=(0, 0, 0), font=small)
    d.text((240, 232), LAN1_ADDR, fill=(0, 0, 0), font=small)

    d.line([520, 230, 700, 230, 700, 300], fill=(40, 40, 40), width=2)
    d.text((540, 200), "eth2  LAN2", fill=(0, 0, 0), font=small)
    d.text((540, 216), "172.22.32.254/24", fill=(0, 0, 0), font=small)
    d.text((540, 232), LAN2_ADDR, fill=(0, 0, 0), font=small)

    d.text((330, 20), "ITS 3100 Deliverable 01 - IPv6 topology",
           fill=(0, 0, 0), font=font)
    img.save(path)


# The answer key. Derived from what each fixture was written to contain, and
# from the rating rule in judge.rate: Complete needs every Evidence bullet,
# Acceptable needs at least a third, Missing is below that.
EXPECTED = {
    "strong": {
        "Task 2 - IPv4 Deployment and Connectivity": "Complete",
        "Task 3 - SLAAC and DHCPv6 on the WAN": "Complete",
        "Task 4 - DHCP-PD on the LAN Interfaces": "Complete",
        "Task 5 - IPv6 and DNS": "Complete",
        "Task 6 - Router Advertisements on LAN1": "Complete",
        "Task 7 - IPv6-Only Network and NAT64": "Complete",
        "Additional Question 1 - Source Address Selection": "Complete",
        "Additional Question 2 - DHCP-PD Workflow": "Complete",
    },
    "thin": {
        "Task 2 - IPv4 Deployment and Connectivity": "Acceptable",
        "Task 3 - SLAAC and DHCPv6 on the WAN": "Acceptable",
        "Task 4 - DHCP-PD on the LAN Interfaces": "Acceptable",
        "Task 5 - IPv6 and DNS": "Acceptable",
        "Task 6 - Router Advertisements on LAN1": "Acceptable",
        "Task 7 - IPv6-Only Network and NAT64": "Acceptable",
        "Additional Question 1 - Source Address Selection": "Missing",
        "Additional Question 2 - DHCP-PD Workflow": "Acceptable",
    },
    "partial": {
        "Task 2 - IPv4 Deployment and Connectivity": "Acceptable",
        "Task 3 - SLAAC and DHCPv6 on the WAN": "Acceptable",
        "Task 4 - DHCP-PD on the LAN Interfaces": "Complete",
        "Task 5 - IPv6 and DNS": "Complete",
        "Task 6 - Router Advertisements on LAN1": "Acceptable",
        "Task 7 - IPv6-Only Network and NAT64": "Acceptable",
        "Additional Question 1 - Source Address Selection": "Missing",
        "Additional Question 2 - DHCP-PD Workflow": "Complete",
    },
    "screenshots": {
        # With the vision stage run, this should look like "strong". Without it,
        # everything resting on command output collapses. Judging this fixture
        # twice, once with vision and once without, is what measures the value
        # of the vision stage.
        "Task 2 - IPv4 Deployment and Connectivity": "Complete",
        "Task 3 - SLAAC and DHCPv6 on the WAN": "Complete",
        "Task 4 - DHCP-PD on the LAN Interfaces": "Complete",
        "Task 5 - IPv6 and DNS": "Complete",
        "Task 6 - Router Advertisements on LAN1": "Complete",
        "Task 7 - IPv6-Only Network and NAT64": "Complete",
        "Additional Question 1 - Source Address Selection": "Complete",
        "Additional Question 2 - DHCP-PD Workflow": "Complete",
    },
    "missing": {
        "Task 2 - IPv4 Deployment and Connectivity": "Acceptable",
        "Task 3 - SLAAC and DHCPv6 on the WAN": "Complete",
        "Task 4 - DHCP-PD on the LAN Interfaces": "Acceptable",
        "Task 5 - IPv6 and DNS": "Missing",
        "Task 6 - Router Advertisements on LAN1": "Missing",
        "Task 7 - IPv6-Only Network and NAT64": "Missing",
        "Additional Question 1 - Source Address Selection": "Missing",
        "Additional Question 2 - DHCP-PD Workflow": "Acceptable",
    },
}

# The PDF variants carry the same content as their .docx twins, so they carry
# the same answer key. A difference in score between a pair is a defect in the
# PDF path, not a difference in the work.
EXPECTED["strong-pdf"] = dict(EXPECTED["strong"])
EXPECTED["partial-pdf"] = dict(EXPECTED["partial"])

BUILDERS = {
    "strong": build_strong,
    "thin": build_thin,
    "partial": build_partial,
    "screenshots": build_screenshots,
    "missing": build_missing,
}


def main() -> None:
    subs = RUN / "submissions"
    subs.mkdir(parents=True, exist_ok=True)
    imgdir = HERE / "images"
    imgdir.mkdir(parents=True, exist_ok=True)
    make_diagram(imgdir / "diagram.png")

    for name, builder in BUILDERS.items():
        doc = Document()
        builder(doc, imgdir)
        path = subs / f"{name}.docx"
        doc.save(str(path))
        print(f"wrote {path.relative_to(RUN.parent.parent)}")

    make_pdfs()

    (RUN / "expected.json").write_text(
        json.dumps(EXPECTED, indent=2), encoding="utf-8")
    print(f"wrote {(RUN / 'expected.json').relative_to(RUN.parent.parent)}")

    # The card and lab are copied in rather than referenced, so a run directory
    # is self contained and a later rerun grades against the card as it was.
    for src, dest in (
        (HERE.parent / "cards" / "3100-d01.md", RUN / "card.md"),
        (HERE.parent / "labs" / "3100-d01.md", RUN / "lab.md"),
    ):
        dest.write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
        print(f"copied {src.name} to {dest.relative_to(RUN.parent.parent)}")



# --- PDF variants -------------------------------------------------------------
# ITS 3100 Deliverable 01 asks for a PDF, and 11 of the 32 Fall 2026 submissions
# arrived as one. A PDF loses the font metadata python-docx relies on, so
# monospace detection and image extraction take a different path entirely and
# need their own fixture rather than an assumption that it works.

def build_pdf(name: str, builder) -> None:
    """Render a fixture as a PDF, mono blocks in Courier as a real report has."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas as rl_canvas

    path = RUN / "submissions" / f"{name}.pdf"
    c = rl_canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    margin, y = 0.9 * inch, height - 0.9 * inch
    leading = 12

    def newpage():
        nonlocal y
        # Page furniture, so the extractor's furniture stripping is exercised
        # rather than assumed. A real report carries a running head like this.
        c.setFont("Helvetica-Oblique", 8)
        c.drawString(margin, 0.5 * inch, f"ITS 3100 D01 - {name} - page {c.getPageNumber()}")
        c.showPage()
        y = height - 0.9 * inch

    def write(text, mono=False, size=10):
        nonlocal y
        font = "Courier" if mono else "Helvetica"
        c.setFont(font, 8 if mono else size)
        for line in text.split("\n"):
            # Wrap prose; never wrap monospace, since a wrapped command line
            # stops being quotable evidence.
            chunks = [line] if mono else _wrap(line, 95)
            for chunk in chunks:
                if y < 1.1 * inch:
                    newpage()
                c.drawString(margin, y, chunk)
                y -= leading if not mono else 9.5
        y -= 6

    def image(path_, w=4.6 * inch):
        nonlocal y
        from PIL import Image as PILImage
        with PILImage.open(path_) as im:
            ratio = im.size[1] / im.size[0]
        h = w * ratio
        if y - h < 1.1 * inch:
            newpage()
        c.drawImage(str(path_), margin, y - h, width=w, height=h)
        y -= h + 10

    # Reuse the same content as the .docx builders by recording their calls.
    for kind, payload in _script(builder):
        if kind == "heading":
            write(payload, size=13)
        elif kind == "para":
            write(payload)
        elif kind == "block":
            write(payload, mono=True)
        elif kind == "image":
            image(payload)
    newpage()
    c.save()
    print(f"wrote {path.relative_to(RUN.parent.parent)}")


def _wrap(line: str, width: int) -> list[str]:
    words, out, cur = line.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            out.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    out.append(cur)
    return out or [""]


class _Recorder:
    """Stands in for a docx Document so one builder can emit either format."""

    def __init__(self):
        self.calls: list[tuple[str, object]] = []

    def add_heading(self, text, level=1):
        self.calls.append(("heading", text))

    def add_paragraph(self, style=None):
        self.calls.append(("para", ""))
        return _Run(self)

    def add_picture(self, path):
        self.calls.append(("image", path))


class _Run:
    def __init__(self, rec): self.rec = rec
    def add_run(self, text):
        self.rec.calls[-1] = ("para", text)
        return self
    @property
    def font(self): return self
    name = None
    size = None


def _script(builder):
    """Replay a docx builder against a recorder to get its content sequence."""
    rec = _Recorder()
    imgdir = HERE / "images"
    builder(rec, imgdir)
    merged, buf = [], []
    for kind, payload in rec.calls:
        if kind == "para" and isinstance(payload, str) and "\n" not in payload \
           and payload.startswith(("vyos@", "brandon@", "C:\\", "No.", "  ", "Codes:",
                                   "PING", "64 bytes", "---", "set ", ";;", "Frame",
                                   "Interface", "eth", "S>*", "C>*", "Capture", "|")):
            buf.append(payload)
            continue
        if buf:
            merged.append(("block", "\n".join(buf)))
            buf = []
        if payload:
            merged.append((kind, payload))
    if buf:
        merged.append(("block", "\n".join(buf)))
    return merged


def make_pdfs() -> None:
    build_pdf("strong-pdf", build_strong)
    build_pdf("partial-pdf", build_partial)

if __name__ == "__main__":
    main()
